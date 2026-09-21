"""I/O boundary adapters (REQUIREMENT.md §12, §16; scoreboard G2/G3/G9/G10).

Loaders turn an input port's value/asset into item(s) {json, assets}; the models
only ever see the `json` projection. Materializers turn an output port's result
JSON into an artifact (or value/action). Bytes live in the asset store (§15).

Input types:  text · json · table · document · image · audio
Output types: text · json · table_enrich · document · image · audio · action
"""
import csv
import io

from . import assets, logsetup, tools

log = logsetup.get("adapters")

XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

# Registries advertised in the capability catalog + validated at task creation.
INPUT_TYPES = {
    "text":     {"doc": "a single text string"},
    "json":     {"doc": "a JSON record"},
    "table":    {"doc": "a CSV/Excel file, loaded whole as {rows, columns} for one loop"},
    "document": {"doc": "a PDF/DOCX/Markdown/HTML/TXT file → extracted text"},
    "image":    {"doc": "an image file → caption + OCR text (via vision tool)"},
    "audio":    {"doc": "an audio file → transcript (via transcription tool)"},
}
OUTPUT_TYPES = {
    "text":         {"doc": "a text response"},
    "json":         {"doc": "a structured JSON record"},
    "table_enrich": {"doc": "the input table with new columns merged in"},
    "document":     {"doc": "a rendered PDF/DOCX from Markdown"},
    "image":        {"doc": "an image artifact (from an image_gen tool node)"},
    "audio":        {"doc": "an audio artifact (from a tts tool node)"},
    "action":       {"doc": "an allowlisted side effect + a record of it"},
}
ASSET_INPUT_TYPES = {"table", "document", "image", "audio"}  # value is an uploaded asset


def item(json_value, assets_map=None):
    return {"json": json_value, "assets": assets_map or {}}


# ============================ INGEST (loaders) ============================
def load_port(port, value):
    """Load one input port into a single item {json, assets}. `value` is a raw
    value (text/json) or an asset ref (table/document/image/audio). ONE bundle =
    one loop: a file input is loaded WHOLE (never exploded into multiple loops)."""
    t = port["type"]
    if t == "text":
        return item(value if isinstance(value, str) else str(value))
    if t == "json":
        return item(value)
    if t == "table":
        rows, cols = _read_table(value)   # whole sheet as one input value
        return item({"rows": rows, "columns": cols, "asset": value}, {"source": value})
    if t == "document":
        return item({"text": _read_document(value), "asset": value}, {"source": value})
    if t == "image":
        proj = tools.run("vision", {"asset": value})
        return item({"caption": proj.get("text", ""), "asset": value}, {"source": value})
    if t == "audio":
        proj = tools.run("transcribe", {"asset": value})
        return item({"transcript": proj.get("text", ""), "asset": value}, {"source": value})
    raise ValueError(f"unknown input type '{t}'")


def _read_table(asset_ref):
    data = assets.get_bytes(asset_ref["asset_id"])
    name = (asset_ref.get("name") or "").lower()
    if name.endswith(".csv") or asset_ref.get("mime") == "text/csv":
        text = data.decode("utf-8-sig", "replace")
        reader = csv.DictReader(io.StringIO(text))
        rows = [dict(r) for r in reader]
        return rows, list(reader.fieldnames or [])
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    it = ws.iter_rows(values_only=True)
    header = [str(h) if h is not None else f"col{i}" for i, h in enumerate(next(it, []))]
    rows = []
    for raw in it:
        if raw is None or all(c is None for c in raw):
            continue
        rows.append({header[i]: ("" if c is None else c) for i, c in enumerate(raw) if i < len(header)})
    wb.close()
    return rows, header


def _read_document(asset_ref):
    data = assets.get_bytes(asset_ref["asset_id"])
    name = (asset_ref.get("name") or "").lower()
    if name.endswith(".pdf") or asset_ref.get("mime") == "application/pdf":
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(data))
        return "\n\n".join((p.extract_text() or "") for p in reader.pages).strip()
    if name.endswith(".docx"):
        import docx
        doc = docx.Document(io.BytesIO(data))
        return "\n".join(p.text for p in doc.paragraphs).strip()
    text = data.decode("utf-8", "replace")
    if name.endswith((".html", ".htm")):
        from bs4 import BeautifulSoup
        return BeautifulSoup(text, "html.parser").get_text("\n").strip()
    return text.strip()  # md / txt / anything text


# ============================ EGRESS (materializers) ============================
def materialize_port(port, value):
    """Turn one output port's result value into the final artifact/value.
    Returns {"value": <json>} and/or {"asset": <ref>}."""
    t = port["type"]
    if t in ("text", "json"):
        return {"value": value}
    if t == "document":
        if isinstance(value, dict) and value.get("asset"):
            return {"asset": value["asset"], "value": value}
        md = value.get("markdown") or value.get("text") if isinstance(value, dict) else str(value)
        res = tools.run("render_doc", {"markdown": md or "", "format": port.get("settings", {}).get("format", "pdf"),
                                       "name": port["name"]})
        return {"asset": res["asset"], "value": value}
    if t in ("image", "audio"):
        ref = value.get("asset") if isinstance(value, dict) else None
        if not ref:
            raise ValueError(f"output port '{port['name']}' ({t}) expected an asset (from an image_gen/tts tool node)")
        return {"asset": ref, "value": value}
    if t == "table_enrich":
        # One bundle = one loop: the tree emits the full result rows; we write them
        # to a new file. Accepts {"rows":[{...}]} or a bare list of row objects.
        rows = value.get("rows") if isinstance(value, dict) else value
        if not isinstance(rows, list) or not all(isinstance(r, dict) for r in rows):
            raise ValueError(f"output port '{port['name']}' (table_enrich) expects a list of row objects "
                             "or {\"rows\": [ {...}, ... ]}")
        return {"asset": write_table(rows, port["name"]), "value": {"row_count": len(rows)}}
    if t == "action":
        return {"value": value}  # the action itself is performed by a tool node; this records it
    raise ValueError(f"unknown output type '{t}'")


def write_table(rows, name="table"):
    """Write a list of row dicts to a new Excel asset. Returns an asset ref."""
    cols = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(cols)
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    b = io.BytesIO(); wb.save(b)
    ref = assets.put(b.getvalue(), name=f"{name}.xlsx", mime=XLSX_MIME, meta={"rows": len(rows), "columns": cols})
    log.info("table written: %d row(s), %d column(s) -> asset %s", len(rows), len(cols), ref["asset_id"])
    return ref
